"""Build an Excel report from the shared train/validation/test experiment."""

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
RESULTS_PATH = ROOT / "tmp" / "ml_comparison_results.json"
OUTPUT_PATH = ROOT / "outputs" / "ml-comparison-shared-split" / "ml_algorithm_results.xlsx"


def add_sheet(workbook, name, headers, rows):
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for column in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 34)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    if rows:
        table = Table(displayName=name.replace(" ", "") + "Table", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    return sheet


def build(result):
    workbook = Workbook()
    workbook.remove(workbook.active)
    datasets = [item["dataset"] for item in result["datasets"]]
    ranked = []
    for dataset in datasets:
        models = sorted(
            (item for item in result["model_results"] if item["dataset"] == dataset),
            key=lambda item: item["validation"]["mae"],
        )
        ranked.extend((dataset, rank, item) for rank, item in enumerate(models, 1))
    add_sheet(workbook, "Summary",
        ["Dataset", "Validation rank", "Algorithm", "Train rows", "Validation rows", "Test rows",
         "Validation MAE (s)", "Test MAE (s)", "Validation RMSE (s)", "Test RMSE (s)",
         "Validation R²", "Test R²"],
        [[dataset, rank, item["algorithm"], item["train_rows"], item["validation_rows"],
          item["test_rows"], item["validation"]["mae"], item["test"]["mae"],
          item["validation"]["rmse"], item["test"]["rmse"],
          item["validation"]["r2"], item["test"]["r2"]]
         for dataset, rank, item in ranked])
    add_sheet(workbook, "All Metrics",
        ["Dataset", "Algorithm", "Split", "MAE (s)", "MSE (s²)", "RMSE (s)", "R²", "Fit time (s)"],
        [[item["dataset"], item["algorithm"], split, item[split]["mae"], item[split]["mse"],
          item[split]["rmse"], item[split]["r2"], item["fit_time_s"]]
         for item in result["model_results"] for split in ("validation", "test")])
    add_sheet(workbook, "XGB by Liquid",
        ["Dataset", "Split", "Liquid", "Rows", "Actual mean (s)", "Prediction mean (s)",
         "MAE (s)", "Overall MAE (s)"],
        [[item[key] for key in ("dataset", "split", "liquid", "count", "actual_mean",
                                "prediction_mean", "mae", "overall_mae")]
         for item in result["xgboost_liquid_summary"]])
    add_sheet(workbook, "XGB Predictions",
        ["Dataset", "Split", "Row", "Liquid", "Height (cm)", "Shape", "Mass (g)",
         "Actual (s)", "Prediction (s)", "Absolute error (s)"],
        [[item[key] for key in ("dataset", "split", "row", "liquid", "height_cm", "shape",
                                "mass_g", "actual_s", "prediction_s", "absolute_error_s")]
         for item in result["xgboost_predictions"]])
    add_sheet(workbook, "Split Rows", ["Dataset", "Row", "Split"],
        [[item["dataset"], item["row"], item["split"]] for item in result["split_rows"]])
    add_sheet(workbook, "Datasets",
        ["Dataset", "Rows", "Columns", "Missing", "Liquids", "Shapes", "Train", "Validation", "Test"],
        [[item[key] for key in ("dataset", "rows", "columns", "missing_values", "liquids",
                                "shapes", "train_rows", "validation_rows", "test_rows")]
         for item in result["datasets"]])
    protocol = result["protocol"]
    add_sheet(workbook, "Configuration", ["Setting", "Value"],
        [["Method", protocol["method"]],
         ["Train fraction", protocol["train_fraction"]],
         ["Validation fraction", protocol["validation_fraction"]],
         ["Test fraction", protocol["test_fraction"]],
         ["Shuffle", str(protocol["shuffle"])],
         ["Random state", protocol["random_state"]],
         ["Model selection", protocol["selection_metric"]],
         ["Fitted on", protocol["fit_rows"]],
         ["Features", ", ".join(result["features"])],
         ["Target", result["target"]],
         ["Generated UTC", result["generated_at_utc"]]])
    return workbook


def main():
    result = json.loads(RESULTS_PATH.read_text(encoding="utf-8"))
    if result.get("protocol", {}).get("method") != "shared_holdout":
        raise ValueError("Run tmp/run_all_algorithms.py first; results use an old protocol")
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    build(result).save(OUTPUT_PATH)
    workbook = load_workbook(OUTPUT_PATH, read_only=True)
    assert workbook["Summary"].max_row == 1 + len(result["model_results"])
    assert workbook["Split Rows"].max_row == 1 + len(result["split_rows"])
    assert workbook["XGB Predictions"].max_row == 1 + len(result["xgboost_predictions"])
    print(f"Saved and verified {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
