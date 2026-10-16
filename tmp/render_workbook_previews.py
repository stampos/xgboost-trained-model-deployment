"""Render quick visual previews of the shared-split workbook."""

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "outputs" / "ml-comparison-shared-split" / "ml_algorithm_results.xlsx"
PREVIEW_DIR = ROOT / "tmp" / "workbook_previews" / "shared_split"


def draw_table(axis, sheet, title, max_rows=11, max_columns=8):
    rows = [
        [f"{cell.value:.4f}" if isinstance(cell.value, float) else str(cell.value or "")
         for cell in row]
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, max_rows),
                                   min_col=1, max_col=min(sheet.max_column, max_columns))
    ]
    axis.axis("off")
    axis.set_title(title, loc="left", color="#1F4E78", fontweight="bold")
    table = axis.table(cellText=rows, cellLoc="left", loc="center")
    table.auto_set_font_size(False)
    table.set_fontsize(7)
    table.scale(1, 1.4)
    for (row, _), cell in table.get_celld().items():
        if row == 0:
            cell.set_facecolor("#1F4E78")
            cell.get_text().set_color("white")


def main():
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    figure, axis = plt.subplots(figsize=(18, 5))
    draw_table(axis, workbook["Summary"], "Models ranked by validation MAE", max_columns=12)
    figure.tight_layout()
    figure.savefig(PREVIEW_DIR / "summary_preview.png", dpi=150, bbox_inches="tight")
    plt.close(figure)

    sheets = ["Summary", "All Metrics", "XGB by Liquid", "XGB Predictions", "Split Rows", "Datasets"]
    figure, axes = plt.subplots(3, 2, figsize=(20, 14))
    for axis, name in zip(axes.flat, sheets):
        draw_table(axis, workbook[name], name, max_rows=7)
    figure.tight_layout()
    figure.savefig(PREVIEW_DIR / "all_sheets_preview.png", dpi=150, bbox_inches="tight")
    plt.close(figure)
    print(PREVIEW_DIR)


if __name__ == "__main__":
    main()
